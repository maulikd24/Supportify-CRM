import csv
import io
import json
from pathlib import Path
import aiosqlite
from collections import Counter, defaultdict
from fastapi import APIRouter, Depends, Request, Form, Query
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from app.database import get_db
from app.assessor import analyze_dsat, parse_manual_conversation
from app.zendesk_client import ZendeskClient
from app.auth import require_user

router = APIRouter(tags=["DSAT"])
templates = Jinja2Templates(directory=str(Path(__file__).parent.parent.parent / "templates"))


# ── helpers ────────────────────────────────────────────────────────────────

def _parse_dsat(row) -> dict:
    d = dict(row)
    json_lists = ("root_causes", "csat_recovery_recommendations", "prevention_tips")
    for field in json_lists:
        d[field] = json.loads(d.get(field) or "[]")
    d["follow_up_response"] = json.loads(d.get("follow_up_response") or "{}")
    d["raw_conversation"] = json.loads(d.get("raw_conversation") or "[]")
    return d


def _flat_categories(root_causes: list) -> list[str]:
    return [rc.get("category", "other") for rc in root_causes if rc.get("category")]


async def _dashboard_data(db, agent_filter: str = "") -> dict:
    """Aggregate DSAT statistics, optionally filtered to one agent."""
    where = "WHERE agent_name = ?" if agent_filter else ""
    params = (agent_filter,) if agent_filter else ()

    cursor = await db.execute(
        f"SELECT * FROM dsat_analyses {where} ORDER BY created_at DESC", params
    )
    rows = [_parse_dsat(r) for r in await cursor.fetchall()]

    total = len(rows)
    prob_counts = Counter(r.get("recovery_probability") or "unknown" for r in rows)

    # Category frequency across all root_causes
    cat_counter: Counter = Counter()
    for r in rows:
        cat_counter.update(_flat_categories(r["root_causes"]))

    # Per-agent summary
    agent_map: dict = defaultdict(lambda: {"total": 0, "categories": Counter(), "probs": Counter()})
    for r in rows:
        name = r.get("agent_name") or "Unknown"
        agent_map[name]["total"] += 1
        agent_map[name]["categories"].update(_flat_categories(r["root_causes"]))
        agent_map[name]["probs"][r.get("recovery_probability") or "unknown"] += 1

    agents = [
        {
            "name": name,
            "total": v["total"],
            "top_category": v["categories"].most_common(1)[0][0] if v["categories"] else "—",
            "high_prob": v["probs"].get("high", 0),
            "medium_prob": v["probs"].get("medium", 0),
            "low_prob": v["probs"].get("low", 0),
        }
        for name, v in sorted(agent_map.items(), key=lambda x: -x[1]["total"])
    ]

    # Distinct agent names for filter dropdown
    cursor2 = await db.execute(
        "SELECT DISTINCT agent_name FROM dsat_analyses WHERE agent_name IS NOT NULL ORDER BY agent_name"
    )
    all_agents = [r[0] for r in await cursor2.fetchall()]

    return {
        "total": total,
        "prob_counts": dict(prob_counts),
        "category_counts": cat_counter.most_common(),
        "agents": agents,
        "all_agents": all_agents,
        "rows": rows,
        "agent_filter": agent_filter,
    }


# ── Routes ─────────────────────────────────────────────────────────────────

@router.get("/dsat", response_class=HTMLResponse)
async def dsat_form(
    request: Request,
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    cursor = await db.execute(
        "SELECT * FROM dsat_analyses ORDER BY created_at DESC LIMIT 20"
    )
    history = [_parse_dsat(r) for r in await cursor.fetchall()]
    return templates.TemplateResponse("dsat.html", {
        "request": request,
        "current_user": current_user,
        "history": history,
        "zendesk_configured": ZendeskClient.is_configured(),
    })


@router.post("/dsat", response_class=HTMLResponse)
async def dsat_submit(
    request: Request,
    input_mode: str = Form("manual"),
    ticket_id: str = Form(""),
    manual_subject: str = Form(""),
    agent_name: str = Form(""),
    customer_name: str = Form(""),
    context: str = Form(""),
    manual_conversation: str = Form(""),
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    error = None
    conversation = []
    subject = manual_subject.strip()
    save_ticket_id = ticket_id.strip()

    try:
        if input_mode == "zendesk" and save_ticket_id:
            zd = ZendeskClient()
            ticket_data = await zd.get_ticket_with_conversation(save_ticket_id)
            conversation = ticket_data["conversation"]
            subject = subject or ticket_data["ticket"].get("subject", "")
            if not customer_name:
                customer_name = ticket_data.get("customer_name", "")
        else:
            raw = manual_conversation.strip()
            if not raw:
                raise ValueError("Please paste a conversation.")
            conversation = parse_manual_conversation(raw)
            if not conversation:
                raise ValueError("Could not parse conversation — please use [CUSTOMER]: / [AGENT]: labels.")

        result = await analyze_dsat(
            conversation  = conversation,
            subject       = subject,
            customer_name = customer_name.strip(),
            context       = context.strip(),
        )

        await db.execute(
            """INSERT INTO dsat_analyses
               (ticket_id, ticket_subject, agent_name, customer_name, context, raw_conversation,
                what_went_wrong, root_causes, customer_impact,
                csat_recovery_recommendations, follow_up_response,
                prevention_tips, recovery_probability, recovery_rationale)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                save_ticket_id or None,
                subject or None,
                agent_name.strip() or None,
                customer_name.strip() or None,
                context or None,
                json.dumps(conversation),
                result.get("what_went_wrong", ""),
                json.dumps(result.get("root_causes", [])),
                result.get("customer_impact", ""),
                json.dumps(result.get("csat_recovery_recommendations", [])),
                json.dumps(result.get("follow_up_response", {})),
                json.dumps(result.get("prevention_tips", [])),
                result.get("recovery_probability", ""),
                result.get("recovery_rationale", ""),
            ),
        )
        await db.commit()
        cursor = await db.execute("SELECT last_insert_rowid()")
        row_id = (await cursor.fetchone())[0]
        return RedirectResponse(f"/dsat/{row_id}", status_code=303)

    except Exception as e:
        error = str(e)

    cursor = await db.execute("SELECT * FROM dsat_analyses ORDER BY created_at DESC LIMIT 20")
    history = [_parse_dsat(r) for r in await cursor.fetchall()]
    return templates.TemplateResponse("dsat.html", {
        "request": request,
        "current_user": current_user,
        "error": error,
        "history": history,
        "zendesk_configured": ZendeskClient.is_configured(),
        "prefill": {
            "ticket_id": ticket_id,
            "manual_subject": manual_subject,
            "agent_name": agent_name,
            "customer_name": customer_name,
            "context": context,
            "manual_conversation": manual_conversation,
        },
    })


# ── Dashboard ──────────────────────────────────────────────────────────────

@router.get("/dsat/dashboard", response_class=HTMLResponse)
async def dsat_dashboard(
    request: Request,
    agent: str = Query(""),
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    data = await _dashboard_data(db, agent_filter=agent)
    return templates.TemplateResponse("dsat_dashboard.html", {"request": request, "current_user": current_user, **data})


# ── Downloads ──────────────────────────────────────────────────────────────

def _build_export_rows(rows: list[dict]) -> tuple[list, list[dict]]:
    headers = [
        "ID", "Date", "Agent Name", "Ticket ID", "Ticket Subject",
        "Customer Name", "DSAT Categories", "Recovery Probability",
        "What Went Wrong", "Customer Impact", "Recovery Rationale",
        "Improvements Comment",
    ]
    out = []
    for r in rows:
        cats = ", ".join(
            c.replace("_", " ").title()
            for c in _flat_categories(r["root_causes"])
        )
        out.append({
            "ID": r["id"],
            "Date": (r.get("created_at") or "")[:10],
            "Agent Name": r.get("agent_name") or "",
            "Ticket ID": r.get("ticket_id") or "",
            "Ticket Subject": r.get("ticket_subject") or "",
            "Customer Name": r.get("customer_name") or "",
            "DSAT Categories": cats,
            "Recovery Probability": (r.get("recovery_probability") or "").title(),
            "What Went Wrong": r.get("what_went_wrong") or "",
            "Customer Impact": r.get("customer_impact") or "",
            "Recovery Rationale": r.get("recovery_rationale") or "",
            "Improvements Comment": r.get("improvements_comment") or "",
        })
    return headers, out


@router.get("/dsat/download/csv")
async def download_csv(
    request: Request,
    agent: str = Query(""),
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    data = await _dashboard_data(db, agent_filter=agent)
    headers, rows = _build_export_rows(data["rows"])

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=headers)
    writer.writeheader()
    writer.writerows(rows)
    buf.seek(0)

    filename = f"dsat_report{'_' + agent if agent else ''}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/dsat/download/excel")
async def download_excel(
    request: Request,
    agent: str = Query(""),
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    data = await _dashboard_data(db, agent_filter=agent)
    headers, rows = _build_export_rows(data["rows"])

    wb = openpyxl.Workbook()

    # ── Sheet 1: All records ──
    ws = wb.active
    ws.title = "DSAT Analyses"

    header_fill = PatternFill("solid", fgColor="DC2626")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row[key])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        # Colour rows by recovery probability
        prob = row.get("Recovery Probability", "").lower()
        fill_color = {"high": "DCFCE7", "medium": "FEF9C3", "low": "FEE2E2"}.get(prob, "FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = PatternFill("solid", fgColor=fill_color)

    # Column widths
    col_widths = [6, 12, 18, 12, 28, 18, 28, 18, 50, 40, 40, 50]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ── Sheet 2: Agent summary ──
    ws2 = wb.create_sheet("Agent Summary")
    ag_headers = ["Agent Name", "Total DSATs", "Top Category", "High Recovery", "Medium Recovery", "Low Recovery"]
    for col, h in enumerate(ag_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, ag in enumerate(data["agents"], 2):
        ws2.cell(row=row_idx, column=1, value=ag["name"])
        ws2.cell(row=row_idx, column=2, value=ag["total"])
        ws2.cell(row=row_idx, column=3, value=ag["top_category"].replace("_", " ").title())
        ws2.cell(row=row_idx, column=4, value=ag["high_prob"])
        ws2.cell(row=row_idx, column=5, value=ag["medium_prob"])
        ws2.cell(row=row_idx, column=6, value=ag["low_prob"])
    for i, w in enumerate([22, 14, 24, 16, 18, 14], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # ── Sheet 3: Category breakdown ──
    ws3 = wb.create_sheet("Category Breakdown")
    ws3.cell(row=1, column=1, value="DSAT Category").fill = header_fill
    ws3.cell(row=1, column=1).font = header_font
    ws3.cell(row=1, column=2, value="Count").fill = header_fill
    ws3.cell(row=1, column=2).font = header_font
    for row_idx, (cat, count) in enumerate(data["category_counts"], 2):
        ws3.cell(row=row_idx, column=1, value=cat.replace("_", " ").title())
        ws3.cell(row=row_idx, column=2, value=count)
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"dsat_report{'_' + agent if agent else ''}.xlsx"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Detail + comment ───────────────────────────────────────────────────────

@router.get("/dsat/{analysis_id}", response_class=HTMLResponse)
async def dsat_detail(
    analysis_id: int,
    request: Request,
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    cursor = await db.execute("SELECT * FROM dsat_analyses WHERE id = ?", (analysis_id,))
    row = await cursor.fetchone()
    if not row:
        return templates.TemplateResponse("error.html", {"request": request, "error": "Analysis not found."}, status_code=404)
    analysis = _parse_dsat(row)
    return templates.TemplateResponse("dsat_detail.html", {"request": request, "current_user": current_user, "analysis": analysis})


@router.post("/dsat/{analysis_id}/comment", response_class=HTMLResponse)
async def save_comment(
    analysis_id: int,
    request: Request,
    improvements_comment: str = Form(""),
    current_user=Depends(require_user),
    db=Depends(get_db),
):
    if isinstance(current_user, RedirectResponse): return current_user
    await db.execute(
        "UPDATE dsat_analyses SET improvements_comment = ? WHERE id = ?",
        (improvements_comment.strip() or None, analysis_id),
    )
    await db.commit()
    return RedirectResponse(f"/dsat/{analysis_id}?saved=1", status_code=303)
